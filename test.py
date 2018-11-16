import unittest
import os
from utils.fbo_nightly_scraper import NightlyFBONotices
from fixtures import nightly_file, json_str, filtered_json_str, nightly_data
from utils.get_fbo_attachments import FboAttachments
from fpdf import FPDF
from docx import Document
from bs4 import BeautifulSoup
import requests
import httpretty


def exceptionCallback(request, uri, headers):
    '''
    Create a callback body that raises an exception when opened. This simulates a bad request.
    '''
    raise requests.ConnectionError('Raising a connection error for the test. You can ignore this!')


class NightlyFBONoticesTestCase(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        notice_types= ['MOD','PRESOL','COMBINE', 'AMDCSS']
        naics = ['334111', '334118', '3343', '33451', '334516', '334614', '5112',
                 '518', '54169', '54121', '5415', '54169', '61142']
        cls.nfbo = NightlyFBONotices(date=20180506, notice_types=notice_types, naics=naics)
        cls.file_lines = nightly_file.nightly_file

    @classmethod
    def tearDownClass(cls):
        cls.nfbo = None
        cls.file_lines = None
        cls.empty_file_lines = None

    def test__id_and_count_notice_tags(self):
        result = NightlyFBONoticesTestCase.nfbo._id_and_count_notice_tags(NightlyFBONoticesTestCase.file_lines)
        expected = {'PRESOL': 1, 'COMBINE': 1, 'ARCHIVE': 1, 
                    'AWARD': 1, 'MOD': 1, 'AMDCSS': 1, 'SRCSGT': 1, 'UNARCHIVE': 1}
        self.assertDictEqual(result, expected)

    def test__merge_dicts(self):
        notice = [{'a': '123'}, {'b': '345'}, {'c': '678'}, {'c': '9'}]
        result = NightlyFBONoticesTestCase.nfbo._merge_dicts(notice)
        expected = {'a': '123', 'b': '345', 'c': '678 9'}
        self.assertDictEqual(result, expected)

    def test_pseudo_xml_to_json(self):
        result = "".join(sorted(NightlyFBONoticesTestCase.nfbo.pseudo_xml_to_json(NightlyFBONoticesTestCase.file_lines)))
        expected = "".join(sorted(json_str.json_str))
        self.assertEqual(result, expected)

    def test_filter_json(self):
        result = "".join(sorted(NightlyFBONoticesTestCase.nfbo.filter_json(json_str.json_str)))
        expected = "".join(sorted(filtered_json_str.filtered_json_str))
        self.assertEqual(result, expected)
    

class FboAttachmentsTestCase(unittest.TestCase):
    
    def setUp(self):
        self.fake_fbo_url = 'https://www.fbo.gov/fake'
        self.fboa = FboAttachments(nightly_data = nightly_data.nightly_data)
        text = "This is a test"
        temp_outfile_path = 'temp_test_file'
        with open(temp_outfile_path, 'w') as f:
            f.write(text)
        self.temp_outfile_path = temp_outfile_path
        
        txt_text = "This is a test"
        temp_outfile_path_txt = 'temp_test_file_txt.txt'
        with open(temp_outfile_path_txt, 'w') as f:
            f.write(txt_text)
        self.temp_outfile_path_txt = temp_outfile_path_txt

        pdf = FPDF()
        pdf.add_page()
        pdf.set_xy(0, 0)
        pdf.set_font('arial', 'B', 13.0)
        pdf.cell(ln=0, h=5.0, align='L', w=0, txt="This is a test", border=0)
        pdf.output('test.pdf', 'F')
        self.temp_outfile_path_pdf = 'test.pdf'

        document = Document()
        document.add_heading("This is a test", 0)
        document.save('test.docx')
        self.temp_outfile_path_docx = 'test.docx'
        

    def tearDown(self):
        self.fboa = None
        self.fake_fbo_url = None
        self.temp_outfile_path = None
        self.temp_outfile_path_txt = None
        self.temp_outfile_path_pdf = None
        self.temp_outfile_path_docx = None
        if os.path.exists(self.temp_outfile_path):  
            os.remove(self.temp_outfile_path)
        if os.path.exists(self.temp_outfile_path_txt):  
            os.remove(self.temp_outfile_path_txt)
        if os.path.exists(self.temp_outfile_path_pdf):
            os.remove(self.temp_outfile_path_pdf)
        if os.path.exists(self.temp_outfile_path_docx):
            os.remove(self.temp_outfile_path_docx)
        
    @httpretty.activate
    def test_get_divs(self):
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>	
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>


                            </div><!-- widget -->
                        
                            </div>
        '''
        httpretty.register_uri(httpretty.GET, uri=self.fake_fbo_url, body=body_with_div, content_type = "text/html")
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = ['expected div in here']
        self.assertEqual(len(result), len(expected))
    
    @httpretty.activate
    def test_get_divs_wrong_url(self):
        httpretty.register_uri(httpretty.GET, uri=self.fake_fbo_url, body='No divs in here')
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
    
    @httpretty.activate
    def test_get_divs_non200_url(self):
        httpretty.register_uri(httpretty.GET, uri=self.fake_fbo_url, status=404)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
        
    @httpretty.activate
    def test_get_divs_connection_error(self):
        httpretty.register_uri(method=httpretty.GET, uri=self.fake_fbo_url, status=200, body=exceptionCallback)
        result = self.fboa.get_divs(self.fake_fbo_url)
        expected = []
        self.assertEqual(len(result), len(expected))
            
    def test_insert_attachments(self):
        text = "This is a test"
        notice = {'a': '1', 'b': '2'}
        with open(self.temp_outfile_path, 'w') as f:
            f.write(text)
        file_list = [(self.temp_outfile_path,self.fake_fbo_url)]
        result = self.fboa.insert_attachments(file_list, notice)
        expected = {'a': '1',
                    'b': '2',
                    'attachments':[{'text':text, 
                                    'url':self.fake_fbo_url,
                                    'prediction':None, 
                                    'decision_boundary':None,
                                    'validation':None,
                                    'trained':False}]
        }
        self.assertDictEqual(result, expected)

    @httpretty.activate
    def test_size_check(self):
        httpretty.register_uri(httpretty.HEAD, uri=self.fake_fbo_url, body='This is less than 500MB')
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = True
        self.assertEqual(result, expected)

    @httpretty.activate
    def test_size_check_non200_url(self):
        httpretty.register_uri(httpretty.HEAD, uri=self.fake_fbo_url, status=404)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)

    @httpretty.activate
    def test_size_check_connection_error(self):
        httpretty.register_uri(method=httpretty.HEAD, uri=self.fake_fbo_url, status=200, body=exceptionCallback)
        result = self.fboa.size_check(self.fake_fbo_url)
        expected = False
        self.assertEqual(result, expected)
    
    def test_get_filename_from_cd(self):
        content_disposition = 'attachment; filename="test.doc"'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = 'test.doc'
        self.assertEqual(result, expected)

    def test_get_filename_from_bad_cd(self):
        content_disposition = 'no file name in here!'
        result = self.fboa.get_filename_from_cd(content_disposition)
        expected = None
        self.assertEqual(result, expected)
    
    def test_get_file_name(self):
        attachment_url = 'https://test.pdf'
        content_type = 'application/pdf'
        result = self.fboa.get_file_name(attachment_url, content_type)
        expected = 'test.pdf'
        self.assertEqual(result, expected)
    
    def test_get_attachment_text_txt(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_txt)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_pdf(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_pdf)
        expected = "This is a test"
        self.assertEqual(result, expected)

    def test_get_attachment_text_docx(self):
        result = self.fboa.get_attachment_text(self.temp_outfile_path_docx)
        expected = "This is a test"
        self.assertEqual(result, expected)
    
    def test_get_attachment_url_from_div(self):
        div = '<a href="/utils/view?id=798e26de983ca76f9075de687047445a"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result = self.fboa.get_attachment_url_from_div(div)
        expected = 'https://www.fbo.gov/utils/view?id=798e26de983ca76f9075de687047445a'
        self.assertEqual(result, expected)
    
    def test_get_attachment_url_from_div_space(self):
        div = '<a href="http://  https://www.thisisalinktoanattachment.docx"\
               target="_blank" title="Download/View FD2060-17-33119_FORM_158_00.pdf"\
               class="file">FD2060-17-33119_FORM_158_00.pdf</a>'
        div = BeautifulSoup(div, "html.parser")
        result = self.fboa.get_attachment_url_from_div(div)
        expected = 'https://www.thisisalinktoanattachment.docx'
        self.assertEqual(result, expected)
    
    def test_write_attachments(self):
        #TODO don't rely on the url in the div below continuing to exist
        body_with_div = b'''
                            <div class="notice_attachment_ro notice_attachment_last">
                            <div>	
                        <div class="file"><input type="hidden" name="dnf_class_values[procurement_notice_archive][packages][5][files][0][file][0][preview]" value="&lt;a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'&gt;FA852618Q0033_______0001.pdf&lt;/a&gt;"><a href='/utils/view?id=d95550fb782f53357ed65db571ef9186' target='_blank' title='Download/View FA852618Q0033_______0001.pdf' class='file'>FA852618Q0033_______0001.pdf</a> (16.54 Kb)</div>
                        </div>
                            <div><span class="label">Description:</span> Amendment 0001</div>	</div>


                            </div><!-- widget -->
                        
                            </div>
        '''
        soup = BeautifulSoup(body_with_div, "html.parser")
        attachment_divs = soup.find_all('div', {"class": "notice_attachment_ro"})
        result = self.fboa.write_attachments(attachment_divs)
        expected = [('attachments/FA852618Q0033_______0001.pdf', 
                     'https://www.fbo.gov/utils/view?id=d95550fb782f53357ed65db571ef9186')]
        for file_url_tup in result:
            file, _ = file_url_tup
            os.remove(file)
        self.assertEqual(result, expected)
        
        
if __name__ == '__main__':
    unittest.main()